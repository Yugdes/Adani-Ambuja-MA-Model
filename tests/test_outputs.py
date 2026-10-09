"""
Cross-output consistency: every published artefact must equal the model.

Run `python run_model.py` after changing anything; these tests fail if the
summary, dashboard data, memo, README or workbook are stale, if a number is
typed into the dashboard by hand, or if retired figures reappear.
"""

import json
import os
import re
import shutil
import subprocess
import zipfile

import pytest

from model.markdown_generator import build_memo, build_readme


def _read(root, *parts):
    with open(os.path.join(root, *parts), encoding="utf-8") as f:
        return f.read()


def _load_model_data(root):
    text = _read(root, "dashboard", "model_data.js")
    return json.loads(text[text.index("=") + 1:].strip().rstrip(";"))


# ---------------------------------------------------------------- data files

def test_summary_json_matches_model(root, summary):
    assert json.loads(_read(root, "output", "summary.json")) == json.loads(json.dumps(summary))


def test_dashboard_data_matches_model(root, summary):
    assert _load_model_data(root) == json.loads(json.dumps(summary))


def test_memo_is_generated_from_model(root, summary):
    assert _read(root, "deal_rationale", "deal_memo.md") == build_memo(summary)


def test_readme_is_generated_from_model(root, summary):
    assert _read(root, "README.md") == build_readme(summary)


# ---------------------------------------------------------------- no hand-typed numbers in the dashboard

FIGURE = re.compile(r"₹\s?\d|\d+(\.\d+)?\s?%|\b\d{1,3}(,\d{3})+\b|\b\d+(\.\d+)?x\b|\$\d|\b\d+\.\d+\b")


def test_index_html_contains_no_figures(root):
    html = _read(root, "dashboard", "index.html")
    body = html[html.index("<body"):]
    body = re.sub(r"<script.*?</script>|<!--.*?-->", "", body, flags=re.S)
    text = re.sub(r"<[^>]+>", " ", body)
    found = [m.group(0) for m in FIGURE.finditer(text)]
    assert not found, f"hard-coded figures in index.html: {found}"


def test_app_js_contains_no_deal_figures(root):
    js = _read(root, "dashboard", "app.js")
    js = re.sub(r"//.*|/\*.*?\*/", "", js, flags=re.S)
    literals = {float(x) for x in re.findall(r"(?<![\w.$#-])(\d+(?:\.\d+)?)(?![\w%])", js)}
    allowed = {0, 1, 2, 3, 4, 6, 9, 100, 180, 1.1}   # formatting decimals, bar scaling, chart heights
    assert literals <= allowed, f"unexpected numeric literals in app.js: {sorted(literals - allowed)}"


def test_dashboard_renders_without_errors(root, tmp_path):
    node = shutil.which("node")
    if not node:
        pytest.skip("node not installed")
    script = tmp_path / "render.js"
    script.write_text(RENDER_HARNESS, encoding="utf-8")
    out = subprocess.run([node, str(script), os.path.join(root, "dashboard")], capture_output=True, text=True, encoding="utf-8")
    assert out.returncode == 0, out.stdout + out.stderr
    result = json.loads(out.stdout)
    assert result["error"] is None, result["error"]
    assert not result["empty"], f"sections left empty: {result['empty']}"
    assert not result["bad_tokens"], f"'undefined'/'NaN' rendered in: {result['bad_tokens']}"


RENDER_HARNESS = r"""
const fs = require('fs'), path = require('path'), vm = require('vm');
const dir = process.argv[2];
const html = fs.readFileSync(path.join(dir, 'index.html'), 'utf8');
const ids = [...html.matchAll(/id="([^"]+)"/g)].map(m => m[1]);
const els = {};
const mk = id => ({ id, innerHTML: '', textContent: '', hidden: true, dataset: {},
    classList: { add() {}, remove() {} }, addEventListener() {} });
ids.forEach(id => { els[id] = mk(id); });
let ready;
const ctx = { console, Math, Number, String, JSON, Object, Array,
    document: { addEventListener: (e, f) => { ready = f; }, getElementById: id => els[id] || null,
                querySelectorAll: () => [], querySelector: () => null } };
ctx.window = ctx;
vm.createContext(ctx);
vm.runInContext(fs.readFileSync(path.join(dir, 'model_data.js'), 'utf8'), ctx);
vm.runInContext(fs.readFileSync(path.join(dir, 'app.js'), 'utf8'), ctx);
ready();
const containers = ids.filter(id => !['navTabs', 'loadError'].includes(id) && !id.startsWith('tab-'));
const content = id => (els[id].innerHTML || els[id].textContent || '');
const empty = containers.filter(id => !content(id).trim());
const bad = containers.filter(id => /undefined|NaN/.test(content(id)));
console.log(JSON.stringify({ error: els.loadError.hidden === false ? els.loadError.textContent : null, empty, bad_tokens: bad }));
"""


# ---------------------------------------------------------------- retired figures and wording

BANNED = ["70,739", "59.1%", "59.5%", "28,426", "56,700", "8,200", "2,530", "+8.4%", "81,361", "10.5 billion",
          "13.2%", "14.3x", "22,389", "23,094", "55,928", "1,240", "EPS accretion", "accretion/dilution",
          "Bloomberg", "Mergermarket", "Analyst Name", "CONFIDENTIAL", "Interactive", "SBI consortium"]


def _xlsx_text(path):
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))


@pytest.mark.parametrize("parts", [("README.md",), ("deal_rationale", "deal_memo.md"), ("dashboard", "index.html"),
                                   ("dashboard", "app.js"), ("output", "Adani_Ambuja_MA_Model.xlsx")])
def test_no_retired_figures_or_wording(root, parts):
    path = os.path.join(root, *parts)
    text = _xlsx_text(path) if path.endswith(".xlsx") else _read(root, *parts)
    hits = [b for b in BANNED if b.lower() in text.lower()]
    assert not hits, f"{'/'.join(parts)} contains retired content: {hits}"


# ---------------------------------------------------------------- Excel workbook

XLSX = ("output", "Adani_Ambuja_MA_Model.xlsx")


def test_workbook_is_formula_driven(root):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(root, *XLSX))
    for name in ["Deal & Financing", "PPA & Goodwill", "Projections", "Synergies", "Pro Forma & Returns", "Balance Sheet", "Checks"]:
        cells = [c.value for row in wb[name].iter_rows() for c in row if c.value is not None]
        formulas = [v for v in cells if isinstance(v, str) and v.startswith("=")]
        assert len(formulas) >= 10, f"{name} has only {len(formulas)} formulas"
    assert "STATIC OUTPUT FROM PYTHON" in wb["Sensitivity"]["B1"].value


def test_workbook_has_no_numbers_stored_as_text(root):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(root, *XLSX))
    pattern = re.compile(r"^[+\-−]?[₹$]?[\d,]+(\.\d+)?[%x]?$")
    bad = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and pattern.match(c.value.strip()) and not ws.title == "Precedents"]
    assert not bad, f"numbers stored as text: {bad[:10]}"


def test_workbook_recalculates_and_all_checks_pass(root, tmp_path, summary):
    formulas = pytest.importorskip("formulas")
    src = tmp_path / "wb.xlsx"
    shutil.copy(os.path.join(root, *XLSX), src)
    sol = formulas.ExcelModel().loads(str(src)).finish().calculate()
    vals = {str(k).upper().split("]", 1)[1]: v.value[0][0] for k, v in sol.items() if hasattr(v, "value")}
    assert vals["CHECKS'!C3"] == "ALL CHECKS OK"
    statuses = [v for k, v in vals.items() if k.startswith("CHECKS'!") and v in ("OK", "FAIL", "DIFF")]
    assert len(statuses) > 200 and set(statuses) == {"OK"}


def test_workbook_opens_with_values_without_recalculation(root, summary):
    """Excel's Protected View does not recalculate, so every formula must carry a cached result."""
    import openpyxl
    path = os.path.join(root, *XLSX)
    values, formulas_wb = openpyxl.load_workbook(path, data_only=True), openpyxl.load_workbook(path)
    missing = [(ws.title, c.coordinate) for ws in formulas_wb for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith("=") and values[ws.title][c.coordinate].value is None]
    assert not missing, f"formula cells without cached values: {missing[:10]}"
    assert values["Checks"]["C3"].value == "ALL CHECKS OK"
